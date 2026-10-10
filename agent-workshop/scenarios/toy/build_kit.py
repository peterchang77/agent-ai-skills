# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.1.2", "reportlab==4.2.5"]
# ///
"""Regenerate the synthetic 'toy' data kit (deterministic).

    uv run build_kit.py

Writes kit/ (what participants receive) and reference/truth.json (the planted truth,
kept OUT of kit/ so participants and agents do not see it).

Everything here is synthetic. Effects are planted so analyses can be checked against known
answers; the numbers are not real findings.
"""
import csv, json, math, random, sqlite3
from datetime import date, timedelta
from pathlib import Path

HERE = Path(__file__).parent
KIT = HERE / "kit"; KIT.mkdir(exist_ok=True)
REF = HERE / "reference"; REF.mkdir(exist_ok=True)
rng = random.Random(2026)

N_ADULTS, N_MINORS = 2400, 8
CLINIC_W = {"Northside": 0.30, "Riverside": 0.25, "Hillcrest": 0.45}
DIAB_MEAN = {"Northside": 7.1, "Riverside": 7.8, "Hillcrest": 7.5}   # latent HbA1c, diabetic pts
PILOT_START, INDEX_END = date(2025, 4, 1), date(2025, 8, 31)
SAT_START = date(2025, 3, 8)
TRUE_LOG_OR_CALL = 0.55          # effect of a reached call on odds of HbA1c recheck
AGE_EFFECT_RECHECK, AGE_EFFECT_CALL = 0.03, 0.04   # confounding: older -> more calls, more rechecks

logistic = lambda x: 1 / (1 + math.exp(-x))
def poisson(lam):
    L, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= L: return k
        k += 1
def rdate(a, b): return a + timedelta(days=rng.randint(0, (b - a).days))

# ---------------- patients and visits ----------------
patients = {}
for i in range(1, N_ADULTS + N_MINORS + 1):
    minor = i > N_ADULTS
    clinic = rng.choices(list(CLINIC_W), weights=list(CLINIC_W.values()))[0]
    age = rng.randint(14, 17) if minor else int(min(92, max(18, rng.gauss(57, 16))))
    diabetic = rng.random() < (0.3 if minor else logistic(-2.4 + 0.045 * age))
    latent = rng.gauss(DIAB_MEAN[clinic], 1.1) if diabetic else rng.gauss(5.4, 0.3)
    patients[f"P{i:04d}"] = dict(clinic=clinic, age=age, sex=rng.choice("FM"),
                                 diabetic=diabetic, latent=min(12.5, max(4.5, latent)))

def make_visit(pid, d, force_type=None):
    p = patients[pid]; c = p["clinic"]
    p_urg = 0.20 + (0.05 if d.month in (1, 2, 11, 12) else 0)
    if c == "Northside" and d >= SAT_START: p_urg -= 0.07          # Saturday-hours effect
    vt = force_type or ("urgent" if rng.random() < p_urg else
                        "follow-up" if rng.random() < (0.4 if p["diabetic"] else 0.15) else "routine")
    p_meas = (0.8 if vt != "urgent" else 0.15) if p["diabetic"] else (0.25 if vt != "urgent" else 0.05)
    a1c = round(min(14, max(4, p["latent"] + rng.gauss(0, 0.35))), 1) if rng.random() < p_meas else None
    return dict(patient_id=pid, visit_date=d, clinic=c, age=p["age"], sex=p["sex"], visit_type=vt,
                systolic_bp=int(rng.gauss(112 + 0.4 * p["age"] + (6 if p["diabetic"] else 0), 12)),
                hba1c=a1c, ldl=None if rng.random() < 0.12 else int(rng.gauss(110, 28)),
                bmi=round(rng.gauss(27 + (2 if p["diabetic"] else 0), 4.5), 1))

visits = []
for pid in patients:
    for _ in range(1 + min(6, poisson(2.0))):
        visits.append(make_visit(pid, rdate(date(2025, 1, 1), date(2025, 12, 31))))

# ---------------- pilot: eligibility, calls, rechecks ----------------
by_pt = {}
for v in visits: by_pt.setdefault(v["patient_id"], []).append(v)
calls, eligible = [], 0
for pid, vs in by_pt.items():
    p = patients[pid]
    if p["clinic"] != "Hillcrest" or p["age"] < 18: continue
    idx = sorted((v for v in vs if v["hba1c"] is not None and v["hba1c"] >= 7.0
                  and PILOT_START <= v["visit_date"] <= INDEX_END), key=lambda v: v["visit_date"])
    if not idx: continue
    index = idx[0]; eligible += 1
    assigned = rng.random() < logistic(-0.3 + AGE_EFFECT_CALL * (p["age"] - 60))
    reached = assigned and rng.random() < 0.9
    if assigned:
        calls.append(dict(patient_id=pid, call_date=index["visit_date"] + timedelta(days=rng.randint(1, 10)),
                          reached="Y" if reached else "N"))
    recheck = rng.random() < logistic(-0.9 + AGE_EFFECT_RECHECK * (p["age"] - 60) + TRUE_LOG_OR_CALL * reached)
    lo, hi = index["visit_date"] + timedelta(days=14), index["visit_date"] + timedelta(days=120)
    for v in vs:                                    # clear any HbA1c in the outcome window
        if lo <= v["visit_date"] <= hi and v is not index: v["hba1c"] = None
    if recheck:
        d = index["visit_date"] + timedelta(days=rng.randint(21, 110))
        nv = make_visit(pid, d, force_type="follow-up")
        nv["hba1c"] = round(min(14, max(4, p["latent"] - 0.3 * reached + rng.gauss(0, 0.35))), 1)
        visits.append(nv)

visits.sort(key=lambda v: (v["visit_date"], v["patient_id"]))
for n, v in enumerate(visits, 1): v["visit_id"] = f"V{n:05d}"
cols = ["visit_id", "patient_id", "visit_date", "clinic", "age", "sex", "visit_type",
        "systolic_bp", "hba1c", "ldl", "bmi"]
rows = [{c: ("" if v[c] is None else v[c]) for c in cols} for v in visits]
for r in rows: r["visit_date"] = r["visit_date"].isoformat()

# ---------------- planted data-quality problems ----------------
non_h = [i for i, r in enumerate(rows) if r["clinic"] != "Hillcrest"]
planted = {}
# 1. clinic-name variants
for i in rng.sample(range(len(rows)), 45):
    rows[i]["clinic"] = {"Northside": rng.choice(["northside ", "NORTHSIDE"]),
                         "Riverside": rng.choice(["riverside", "Riverside "]),
                         "Hillcrest": rng.choice(["hillcrest", "Hillcrest "])}[rows[i]["clinic"].strip()]
# 2. HbA1c reported in mmol/mol instead of % (non-Hillcrest rows only)
cand = [i for i in non_h if rows[i]["hba1c"] != ""]
mm = rng.sample(cand, 10)
for i in mm: rows[i]["hba1c"] = round((float(rows[i]["hba1c"]) - 2.15) * 10.929)
planted["hba1c_mmol_per_mol_rows"] = len(mm)
# 3. one impossible age (patient with other visits at a consistent age)
multi = [i for i in non_h if sum(r["patient_id"] == rows[i]["patient_id"] for r in rows) >= 3]
ai = rng.choice(multi); true_age = rows[ai]["age"]; rows[ai]["age"] = 212
planted["age_212"] = dict(visit_id=rows[ai]["visit_id"], patient_id=rows[ai]["patient_id"], true_age=true_age)
# 4. systolic BP recorded as 0
for i in rng.sample(range(len(rows)), 4): rows[i]["systolic_bp"] = 0
# 5. US-style dates (day > 12 so they are unambiguous)
us = [i for i, r in enumerate(rows) if int(r["visit_date"][8:]) > 12]
for i in rng.sample(us, 6):
    y, m, d = rows[i]["visit_date"].split("-"); rows[i]["visit_date"] = f"{m}/{d}/{y}"
# 6. one conflicting visit_id (two different visits share an id)
a, b = rng.sample(non_h, 2); rows[b]["visit_id"] = rows[a]["visit_id"]
planted["conflicting_id"] = rows[a]["visit_id"]
# 7. exact duplicate rows
dup = sorted(rng.sample(range(len(rows)), 15), reverse=True)
for i in dup: rows.insert(i + 1, dict(rows[i]))
planted["exact_duplicate_rows"] = len(dup)

with open(KIT / "clinic_visits.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(rows)
with open(KIT / "pilot_calls.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["patient_id", "call_date", "reached"]); w.writeheader()
    for c in sorted(calls, key=lambda c: c["patient_id"]): w.writerow({**c, "call_date": c["call_date"].isoformat()})

(REF / "truth.json").write_text(json.dumps(dict(
    true_log_or_call=TRUE_LOG_OR_CALL, age_effect_recheck=AGE_EFFECT_RECHECK, age_effect_call=AGE_EFFECT_CALL,
    saturday_hours_start=SAT_START.isoformat(), pilot_eligible_patients=eligible, planted=planted,
    notes="Synthetic. Call effect is confounded by age (older -> more likely called and more likely to recheck)."),
    indent=2))

# ---------------- text, markdown ----------------
(KIT / "clinic_notes.txt").write_text("""Operations memo: three-clinic network, 2025 (synthetic)

From: Clinical operations lead
To: Quality analyst

Thanks for taking this on. Three things for the quarterly review.

1. TARGETS. Each clinic has a 2025 target for the share of adult patients with diabetes whose
   most recent HbA1c is below 8.0%. The targets are in clinic_targets.xlsx. Which clinics are
   meeting them?

2. HILLCREST PILOT. Since April 1 Hillcrest has run a nurse follow-up call program for
   patients with poorly controlled diabetes (protocol in pilot_protocol.docx). We want to
   know whether patients who were reached by a call were more likely to have their HbA1c
   rechecked in time. Older patients seem to get more calls, so please be careful about
   that. Draft a one-page abstract for the leadership packet, with references to
   related published work.

3. (If time) Northside moved to Saturday-morning hours on March 8. Front-desk staff think
   urgent visits dropped afterward. Is that visible in the data?

Notes on the data. The visit extract came from three different scheduling systems, so
expect some inconsistencies. Please do not edit the original files; put anything you produce
in a separate folder. Clinic hours and staffing context are in clinic_directory.html and
staffing.db.
""")
(KIT / "data_dictionary.md").write_text("""# Data dictionary (synthetic)

## clinic_visits.csv: one row per visit, calendar year 2025

| Column | Meaning | Notes |
|---|---|---|
| visit_id | Visit identifier | `V` + 5 digits; should be unique |
| patient_id | De-identified patient | `P` + 4 digits |
| visit_date | Date of visit | ISO `YYYY-MM-DD` expected |
| clinic | Clinic site | Northside, Riverside, Hillcrest |
| age | Age in years at visit | |
| sex | `F` or `M` | |
| visit_type | routine, follow-up, urgent | |
| systolic_bp | Systolic blood pressure, mmHg | |
| hba1c | Glycated hemoglobin, **percent (NGSP)** | blank if not measured |
| ldl | LDL cholesterol, mg/dL | blank if not measured |
| bmi | Body mass index, kg/m^2 | |

A patient is treated as having diabetes if any HbA1c in 2025 is 6.5% or higher.

## pilot_calls.csv: one row per patient who was called by the Hillcrest pilot

| Column | Meaning |
|---|---|
| patient_id | Matches clinic_visits.csv |
| call_date | Date of the call attempt |
| reached | `Y` if the patient was reached and counseled; `N` if not reached |

## Other files

- `clinic_targets.xlsx`: 2025 targets and staffing (formatted for people, not machines).
- `pilot_protocol.docx`: pilot eligibility and outcome definitions.
- `patient_newsletter.pdf`: patient-facing newsletter.
- `clinic_directory.html`: clinic hours and program announcements.
- `staffing.db`: SQLite database with monthly staffing by clinic.
""")

# ---------------- xlsx ----------------
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = "Targets"
ws["A1"] = "2025 Clinic Targets (draft)"; ws.merge_cells("A1:D1")
ws.append([]); ws.append(["Clinic", "Metric", "Target", "Note"])
for r in [("Northside", "Most recent HbA1c < 8.0% (share of adult pts with diabetes)", 0.60, ""),
          ("Riverside", "Most recent HbA1c < 8.0% (share of adult pts with diabetes)", 0.60, "revisit after staffing"),
          ("Hillcrest", "Most recent HbA1c < 8.0% (share of adult pts with diabetes)", 0.55, "pilot site")]:
    ws.append(list(r))
for c in ws["C"][3:]: c.number_format = "0%"
ws2 = wb.create_sheet("Staffing (planned)")
ws2.append(["Clinic", "Physicians", "NPs", "Nurses"])
for r in [("Northside", 4, 2, 6), ("Riverside", 3, 2, 5), ("Hillcrest", 3, 2, 5)]: ws2.append(list(r))
wb.save(KIT / "clinic_targets.xlsx")

# ---------------- docx ----------------
from docx import Document
d = Document(); d.add_heading("Follow-up Call Pilot: Protocol (synthetic)", 1)
d.add_paragraph("Purpose: test whether a nurse phone call after a visit with poorly controlled "
                "diabetes increases the chance that HbA1c is rechecked on time.")
d.add_heading("Eligibility", 2)
for s in ["Patient is an adult (age 18 or older) seen at Hillcrest.",
          "Index visit: the first visit between 2025-04-01 and 2025-08-31 with a measured HbA1c of 7.0% or higher.",
          "Patients younger than 18 are excluded."]:
    d.add_paragraph(s, style="List Bullet")
d.add_heading("Procedure", 2)
for s in ["A nurse attempts a call within 10 days of the index visit using the standard script.",
          "Record the call date and whether the patient was reached (Y/N) in the tracking sheet."]:
    d.add_paragraph(s, style="List Number")
d.add_heading("Outcome", 2)
d.add_paragraph("HbA1c recheck: any measured HbA1c from 14 to 120 days after the index visit.")
t = d.add_table(rows=2, cols=3); t.style = "Table Grid"
for c, txt in zip(t.rows[0].cells, ["Site", "Start", "Lead"]): c.text = txt
for c, txt in zip(t.rows[1].cells, ["Hillcrest", "2025-04-01", "R. Alvarez, RN"]): c.text = txt
d.save(KIT / "pilot_protocol.docx")

# ---------------- pdf ----------------
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
c = canvas.Canvas(str(KIT / "patient_newsletter.pdf"), pagesize=letter)
c.setFont("Helvetica-Bold", 16); c.drawString(72, 720, "Community Health Newsletter, Fall 2025 (synthetic)")
c.setFont("Helvetica", 11)
for n, line in enumerate(["Saturday hours are now available at our Northside clinic.",
                          "Diabetes? Ask about our free follow-up call program (Hillcrest).",
                          "Flu shots: walk-ins welcome at all three sites through November."]):
    c.drawString(72, 690 - 20 * n, line)
c.save()

# ---------------- html + sqlite (cross-reference sources) ----------------
(KIT / "clinic_directory.html").write_text("""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Clinic directory (synthetic)</title></head>
<body>
<h1>Clinic directory</h1>
<table border="1">
<tr><th>Clinic</th><th>Hours</th><th>Announcements</th></tr>
<tr><td>Northside</td><td>Mon-Fri 8-5; <b>Sat 8-12 (from 2025-03-08)</b></td><td>Saturday hours pilot began March 8, 2025.</td></tr>
<tr><td>Riverside</td><td>Mon-Fri 8-5</td><td>One nurse practitioner position vacant since June 2025.</td></tr>
<tr><td>Hillcrest</td><td>Mon-Fri 8-6</td><td>Diabetes follow-up call program began April 1, 2025.</td></tr>
</table></body></html>
""")
db = KIT / "staffing.db"; db.unlink(missing_ok=True)
con = sqlite3.connect(db)
con.execute("create table staffing (clinic text, month text, physicians integer, nps integer, nurses integer)")
base = {"Northside": (4, 2, 6), "Riverside": (3, 2, 5), "Hillcrest": (3, 2, 5)}
for cl, (p, n, u) in base.items():
    for m in range(1, 13):
        nps = n - 1 if (cl == "Riverside" and m >= 6) else n
        con.execute("insert into staffing values (?,?,?,?,?)", (cl, f"2025-{m:02d}", p, nps, u))
con.commit(); con.close()

# ---------------- free-text notes for bulk classification ----------------
cats = {
 "medication_refill": ["Patient requests refill of {m}. No side effects reported.", "Needs {m} renewed; pharmacy faxed request.", "Out of {m} for two days, asking for early refill."],
 "new_symptom": ["Reports {s} for {d} days, no fever.", "Came in with {s}; started {d} days ago.", "{s} worsening over {d} days, wants evaluation."],
 "chronic_followup": ["Follow-up for diabetes; reviewing glucose log and {m}.", "Hypertension recheck, home BP readings reviewed.", "Routine diabetes visit, discussed diet and {m}."],
 "preventive": ["Here for flu vaccine.", "Annual wellness visit, screening labs ordered.", "Vaccine catch-up and routine screening discussed."],
 "administrative": ["Needs paperwork completed for work.", "Requesting a copy of recent lab results.", "Insurance form to be signed; no clinical complaints."]}
meds, syms = ["metformin", "lisinopril", "atorvastatin", "insulin glargine"], ["cough", "back pain", "dizziness", "rash", "headache"]
nrng = random.Random(7); notes, truth = [], []
for i in range(1, 201):
    cat = nrng.choice(list(cats))
    txt = nrng.choice(cats[cat]).format(m=nrng.choice(meds), s=nrng.choice(syms), d=nrng.randint(2, 14))
    notes.append((f"N{i:03d}", txt)); truth.append((f"N{i:03d}", cat))
with open(KIT / "visit_notes.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["note_id", "note_text"]); w.writerows(notes)
with open(KIT / "visit_notes_labeled_sample.csv", "w", newline="") as f:   # first 30 labels, for estimating accuracy
    w = csv.writer(f); w.writerow(["note_id", "category"]); w.writerows(truth[:30])
with open(REF / "visit_notes_truth.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["note_id", "category"]); w.writerows(truth)
print("kit written:", len(rows), "visit rows;", len(calls), "calls;", eligible, "eligible pilot patients")
