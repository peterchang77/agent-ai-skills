# /// script
# requires-python = ">=3.10"
# dependencies = ["pydicom==2.4.4", "pylibjpeg==2.0.1", "pylibjpeg-libjpeg==2.3.0", "numpy", "pillow==10.4.0", "python-docx==1.1.2"]
# ///
"""Regenerate the imaging workshop data kit (deterministic).

    uv run build_kit.py

Downloads siim_small (250 real, de-identified chest radiographs; cached in .cache/), then writes:
  kit/        what participants receive (DICOM files, labels, model scores, reports, protocol, memo)
  truth/      the planted truth (kept OUT of kit/ so participants and agents do not see it)

The real images are a small subset of the SIIM-ACR Pneumothorax Segmentation challenge data.
Everything else is synthetic: the planted data problems, the fake PHI, the radiology reports, and the
follow-up images. Model scores are precomputed by score_images.py from the PRISTINE images, so a score
never reflects a planted pixel problem (running the model on the uncleaned kit will give different, worse
numbers for those files, which is deliberate).
"""
import copy
import warnings
import csv
import hashlib
import json
import math
import random
import shutil
import sys
import uuid
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pydicom
from PIL import Image, ImageDraw, ImageFont
from pydicom.uid import ExplicitVRLittleEndian

warnings.filterwarnings('ignore', message='Invalid value for VR')
HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from siim_data import fetch_siim_small, sha256  # noqa: E402

KIT, REF = HERE / "kit", HERE / "truth"   # generated; reference/ holds hand-written solution scripts and is never wiped
for d in (KIT, REF):
    shutil.rmtree(d, ignore_errors=True)
(KIT / "images").mkdir(parents=True); REF.mkdir()
rng = random.Random(2026)
nrng = np.random.default_rng(2026)


def uid(tag: str) -> str:
    return "2.25." + str(int(hashlib.sha256(f"imaging-kit:{tag}".encode()).hexdigest(), 16) % 10**30)


def set_pixels(ds, arr):
    arr = np.clip(arr, 0, 255).astype(np.uint8)
    ds.PixelData = arr.tobytes()
    ds.Rows, ds.Columns = arr.shape
    ds.file_meta.TransferSyntaxUID = ExplicitVRLittleEndian
    ds.is_little_endian, ds.is_implicit_VR = True, False


# ---------------- load pristine data ----------------
MODELS = ["all", "chex", "nih", "pc"]
root = fetch_siim_small()
pristine = {}
with open(HERE / "assets" / "pristine_scores.csv") as f:
    for r in csv.DictReader(f):
        pristine[r["source_path"]] = dict(label=r["true_label"], scores={m: float(r[f"score_{m}"]) for m in MODELS})
keys = sorted(pristine)
assert len(keys) == 250

entries = []   # one per kit file
by_src = {}
for k in keys:
    ds = pydicom.dcmread(root / k)
    ds.StudyDate = (date(2025, 1, 1) + timedelta(days=rng.randint(0, 299))).strftime("%Y%m%d")
    e = dict(ds=ds, source=k, kind="original", label=pristine[k]["label"], scores=pristine[k]["scores"], problems=[])
    entries.append(e); by_src[k] = e

pool = keys[:]; rng.shuffle(pool)
take = lambda n: [by_src[pool.pop()] for _ in range(n)]
truth = {}


def plant(name, es, note):
    truth[name] = dict(note=note, n=len(es), source_paths=[e["source"] for e in es])
    for e in es:
        e["problems"].append(name)


# ---------------- pixel problems ----------------
def invert(e, photometric):
    a = e["ds"].pixel_array.astype(np.int16)
    set_pixels(e["ds"], 255 - a); e["ds"].PhotometricInterpretation = photometric


es = take(6)
for e in es: invert(e, "MONOCHROME1")
plant("inverted_tagged_mono1", es, "Pixels inverted and PhotometricInterpretation = MONOCHROME1 (header is honest; the fix is to honor the tag)")
es = take(4)
for e in es: invert(e, "MONOCHROME2")
plant("inverted_tag_wrong", es, "Pixels inverted but tagged MONOCHROME2 (header lies; only looking at the image reveals it)")
es = take(4)
for e in es: set_pixels(e["ds"], np.fliplr(e["ds"].pixel_array))
plant("flipped_left_right", es, "Image mirrored left-right (burned-in markers read backwards). Laterality is unsafe: exclude, do not repair")
es = take(4)
for e in es:
    a = e["ds"].pixel_array
    small = np.asarray(Image.fromarray(a).resize((512, 512), Image.LANCZOS))
    ps = [float(x) * 2 for x in e["ds"].PixelSpacing]
    set_pixels(e["ds"], small); e["ds"].PixelSpacing = ps
plant("low_resolution_512", es, "Matrix 512x512 instead of 1024x1024 (protocol requires 1024); pixel spacing doubled consistently")
corrupt = take(2)
plant("corrupt_truncated", corrupt, "File truncated partway through the pixel data; cannot be decoded")
es = take(5)
fake = [("DOE, JANE", "MRN 4471902", "DOB 03/14/1961"), ("ROE, RICHARD", "MRN 5520318", "DOB 11/02/1958"),
        ("POE, MARY", "MRN 3390027", "DOB 07/21/1974"), ("LOE, SAM", "MRN 6081455", "DOB 01/30/1949"),
        ("MOE, ANN", "MRN 2217760", "DOB 09/09/1966")]
font = ImageFont.load_default(size=30)
for e, lines in zip(es, fake):
    im = Image.fromarray(e["ds"].pixel_array).convert("L"); dr = ImageDraw.Draw(im)
    dr.rectangle([0, 0, 420, 125], fill=0)
    for i, t in enumerate(lines): dr.text((10, 8 + 38 * i), t, fill=255, font=font)
    set_pixels(e["ds"], np.asarray(im))
    e["fake_phi_text"] = list(lines)
plant("burned_in_phi", es, "Synthetic patient name, MRN and birth date drawn into the top-left corner of the pixels (SYNTHETIC; not real people)")

# ---------------- header problems ----------------
es = take(12)
for i, e in enumerate(es):
    p = str(e["ds"].PatientID)
    e["ds"].PatientID = [p.upper(), " " + p, "SITE1-" + p][i // 4]
plant("patient_id_variants", es, "PatientID upper-cased (4), with a leading space (4), or with a SITE1- prefix (4); same patient would not match across files")
es = take(15)
for i, e in enumerate(es):
    s = str(e["ds"].PatientSex)
    e["ds"].PatientSex = [{"M": "MALE", "F": "FEMALE"}[s], "", s.lower()][i // 5]
plant("sex_variants", es, "PatientSex written MALE/FEMALE (5), blank (5), or lower case (5) instead of M/F")
es = take(30)
for i, e in enumerate(es):
    if i < 20: e["ds"].PatientAge = f"{int(e['ds'].PatientAge):03d}Y"
    elif i < 28: e["ds"].PatientAge = ""
    else: e["ds"].PatientAge = ["212Y", "150Y"][i - 28]
truth["age_problems"] = dict(note="Age as 062Y (20; standard DICOM format, harmless; the other 220 files use bare numbers like 62), blank (8), implausible 212Y/150Y (2)",
                             n=30, source_paths=[e["source"] for e in es])
for e in es: e["problems"].append("age_problems")
es = take(10)
for e in es: e["ds"].ViewPosition = ""
plant("view_position_blank", es, "ViewPosition blank, but SeriesDescription still says 'view: PA/AP' (recoverable from another field)")
es = take(6)
for e in es: del e["ds"]["PixelSpacing"]
plant("pixel_spacing_missing", es, "PixelSpacing tag removed")
es = take(5)
for i, e in enumerate(es):
    e["ds"].PatientName = ["DOE^JANE^Q", "ROE^RICHARD", "POE^MARY", "LOE^SAM", "MOE^ANN"][i]
    e["ds"].PatientBirthDate = ["19610314", "19581102", "19740721", "19490130", "19660909"][i]
    e["ds"].OtherPatientIDs = ["MRN4471902", "MRN5520318", "MRN3390027", "MRN6081455", "MRN2217760"][i]
plant("phi_in_header", es, "Synthetic patient name, birth date and MRN in header tags (SYNTHETIC; not real people)")
es = take(10)
for e in es: e["ds"].Modality = "DX"
plant("modality_dx", es, "Modality DX where the rest of the set says CR (both can describe a chest radiograph; the mix is an inconsistency to normalize)")

# ---------------- duplicates, follow-ups, non-chest ----------------
for i, e in enumerate(take(6)):                      # exact duplicates (identical bytes, different file name)
    entries.append(dict(ds=copy.deepcopy(e["ds"]), source=e["source"], kind="exact_duplicate", label=e["label"],
                        scores=e["scores"], problems=["exact_duplicate"], dup_of=e["source"]))
truth["exact_duplicate"] = dict(note="Byte-identical copy of another file under a different name (same SOPInstanceUID)", n=6,
                                source_paths=[x["source"] for x in entries if x["kind"] == "exact_duplicate"])
dup2 = take(4)
for i, e in enumerate(dup2):                          # same pixels, new IDs
    d = copy.deepcopy(e["ds"]); tag = f"redup:{e['source']}"
    d.SOPInstanceUID = d.file_meta.MediaStorageSOPInstanceUID = uid(tag + "sop")
    d.StudyInstanceUID, d.SeriesInstanceUID = uid(tag + "study"), uid(tag + "series")
    d.PatientID = d.PatientName = str(uuid.UUID(int=int(hashlib.sha256(tag.encode()).hexdigest(), 16) % 2**128))
    entries.append(dict(ds=d, source=e["source"], kind="same_pixels_new_ids", label=e["label"], scores=e["scores"],
                        problems=["same_pixels_new_ids"], dup_of=e["source"], conflict_label=(i < 2)))
truth["same_pixels_new_ids"] = dict(note="Same pixels submitted again under new SOP/Study/Series UIDs and a new PatientID. In 2 of the 4 the labels file gives the OPPOSITE label to the original",
                                    n=4, source_paths=[e["source"] for e in dup2])
follow = take(20)
for e in follow:                                      # synthetic follow-up of the same patient
    d = copy.deepcopy(e["ds"]); tag = f"follow:{e['source']}"
    a = d.pixel_array.astype(np.float32)
    a = np.roll(a, (rng.randint(-12, 12), rng.randint(-12, 12)), axis=(0, 1)) * rng.uniform(0.94, 1.06) + nrng.normal(0, 3, a.shape)
    set_pixels(d, a)
    d.SOPInstanceUID = d.file_meta.MediaStorageSOPInstanceUID = uid(tag + "sop")
    d.StudyInstanceUID, d.SeriesInstanceUID = uid(tag + "study"), uid(tag + "series")
    d0 = date(int(str(e["ds"].StudyDate)[:4]), int(str(e["ds"].StudyDate)[4:6]), int(str(e["ds"].StudyDate)[6:8]))
    d.StudyDate = (d0 + timedelta(days=rng.randint(90, 270))).strftime("%Y%m%d")
    sc = {m: float(np.clip(v + rng.gauss(0, 0.03), 0.001, 0.999)) for m, v in e["scores"].items()}
    entries.append(dict(ds=d, source=e["source"], kind="followup", label=e["label"], scores=sc, problems=["followup_same_patient"], dup_of=e["source"]))
truth["followup_same_patient"] = dict(note="Different image (shifted, noise added), same PatientID, later StudyDate, same label: a second study of the same patient. A random image-level split would leak patients across train/test", n=20,
                                      source_paths=[e["source"] for e in follow])
tmpl = by_src[pool.pop()]
for i, (nm, mk) in enumerate([("blank", lambda: nrng.normal(128, 4, (1024, 1024))),
                              ("gradient", lambda: np.tile(np.linspace(20, 235, 1024), (1024, 1))),
                              ("checker", lambda: ((np.indices((1024, 1024)).sum(0) // 64) % 2) * 200 + 25)]):
    d = copy.deepcopy(tmpl["ds"]); tag = f"nonchest:{nm}"
    set_pixels(d, mk())
    d.SOPInstanceUID = d.file_meta.MediaStorageSOPInstanceUID = uid(tag + "sop")
    d.StudyInstanceUID, d.SeriesInstanceUID = uid(tag + "study"), uid(tag + "series")
    d.PatientID = d.PatientName = uid(tag + "pid")[-36:]
    entries.append(dict(ds=d, source=None, kind="non_chest", label=None, scores=None, problems=["non_chest_image"], pattern=nm))
truth["non_chest_image"] = dict(note="Three images that are not chest radiographs (blank noise, gradient, checkerboard) but whose headers say CHEST; no label, no score",
                                n=3, source_paths=[])

# ---------------- write files ----------------
rng.shuffle(entries)
for i, e in enumerate(entries, 1):
    e["name"] = f"img_{i:04d}.dcm"
    e["ds"].save_as(KIT / "images" / e["name"], write_like_original=False)
for e in entries:
    if "corrupt_truncated" in e["problems"]:
        p = KIT / "images" / e["name"]; b = p.read_bytes(); p.write_bytes(b[: int(len(b) * 0.55)])
src2name = {}
for e in entries:
    if e["kind"] == "original": src2name[e["source"]] = e["name"]
for k, v in truth.items():
    v["files"] = sorted({e["name"] for e in entries if k in e["problems"]})

# ---------------- labels.csv ----------------
labeled = [e for e in entries if e["label"] is not None]
rows = {e["name"]: e["label"] for e in labeled}
for e in labeled:
    if e.get("conflict_label"):
        rows[e["name"]] = "No Pneumothorax" if e["label"] == "Pneumothorax" else "Pneumothorax"
plain = [e["name"] for e in labeled if e["kind"] == "original" and not e["problems"]]
rng.shuffle(plain)
variants = {"Pneumothorax": ["PTX", "pneumothorax", "Pneumothorax "], "No Pneumothorax": ["no pneumothorax", "Normal", "NO PNEUMOTHORAX"]}
label_var = []
for lab in variants:
    cand = [n for n in plain if rows[n] == lab][:9]
    for j, n in enumerate(cand):
        rows[n] = variants[lab][j // 3]; label_var.append(n)
out = sorted(rows.items())
out += [("img_9001.dcm", "Pneumothorax"), ("img_9002.dcm", "No Pneumothorax")]
rng.shuffle(out)
with open(KIT / "labels.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["image_file", "label"]); w.writerows(out)
truth["label_spelling_variants"] = dict(note="Label written PTX/pneumothorax/'Pneumothorax ' (9) or 'no pneumothorax'/Normal/'NO PNEUMOTHORAX' (9) instead of the two standard strings",
                                        n=18, files=sorted(label_var))
truth["label_conflict"] = dict(note="For 2 of the 'same_pixels_new_ids' copies, labels.csv gives the opposite label to the original file", n=2,
                               files=sorted(e["name"] for e in entries if e.get("conflict_label")),
                               originals=sorted(src2name[e["dup_of"]] for e in entries if e.get("conflict_label")))
truth["label_without_image"] = dict(note="labels.csv rows for images that are not in images/", n=2, files=["img_9001.dcm", "img_9002.dcm"])

# ---------------- model_scores.csv ----------------
srows = {e["name"]: e["scores"] for e in labeled}
cand = [e["name"] for e in labeled if e["kind"] == "original" and not e["problems"] and e["name"] not in label_var]
rng.shuffle(cand)
# logit rows must be recognizable as such: at least one of the four values falls outside [0, 1]
recog = [n for n in cand if min(srows[n].values()) < 0.5]
logit_f = recog[:25]
rest = [n for n in cand if n not in logit_f]
nan_f, miss_f, case_f = rest[:5], rest[5:11], rest[11:21]
lg = lambda s: f"{math.log(s / (1 - s)):.4f}"
srow_out = []
for n, sc in srows.items():
    if n in miss_f: continue
    vals = [lg(sc[m]) if n in logit_f else ("" if n in nan_f else f"{sc[m]:.6f}") for m in MODELS]
    srow_out.append((n.upper() if n in case_f else n, *vals))
rng.shuffle(srow_out)
with open(KIT / "model_scores.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["image_file", *[f"score_{m}" for m in MODELS]]); w.writerows(srow_out)
truth["scores_as_logits"] = dict(note="25 rows have all four scores as logits (can be negative or above 1) while the rest are probabilities in [0,1]", n=25, files=sorted(logit_f))
truth["scores_blank"] = dict(note="Blank score", n=5, files=sorted(nan_f))
truth["scores_missing_row"] = dict(note="Image has a label but no row in model_scores.csv", n=6, files=sorted(miss_f))
truth["scores_key_case"] = dict(note="image_file written in upper case (IMG_0012.DCM): a case-sensitive join misses them", n=10, files=sorted(case_f))

# ---------------- synthetic radiology reports ----------------
LUNGS = ["The lungs are otherwise clear.", "Mild bibasilar atelectasis.", "Heart size is within normal limits.",
         "Mild cardiomegaly.", "Degenerative changes of the thoracic spine.", "No focal consolidation.", "Small left pleural effusion.",
         "Lung volumes are low.", "Mediastinal contours are unremarkable."]
rr = random.Random(77)
pickable = [e for e in entries if e["kind"] in ("original", "followup") and "corrupt_truncated" not in e["problems"] and e["label"]]
rr.shuffle(pickable)
reports, rtruth = [], []
for i, e in enumerate(pickable[:200], 1):
    pos = e["label"] == "Pneumothorax"
    side = rr.choice(["right", "left"]); view = "Portable AP" if str(e["ds"].SeriesDescription).endswith("AP") else "PA"
    extra = rr.choice(LUNGS)
    r = rr.random()
    if pos and r < 0.80:
        lab = "present"
        txt = rr.choice([f"{view} chest radiograph. {side.capitalize()} apical pneumothorax, {rr.choice(['small', 'moderate', 'large'])}. {extra}",
                         f"{view} view of the chest. There is a {rr.choice(['small', 'moderate'])} {side} pneumothorax. {rr.choice(['Chest tube in place.', 'No tension physiology.', 'Pigtail catheter in place.'])} {extra}",
                         f"Findings: visceral pleural line along the {side} lung with absent lung markings peripherally, consistent with pneumothorax. {extra}"])
    elif pos:
        lab = "uncertain"
        txt = rr.choice([f"{view} chest radiograph. Possible tiny {side} apical pneumothorax; cannot be excluded on this view. Suggest repeat upright film. {extra}",
                         f"Equivocal lucency at the {side} apex, may represent a skin fold versus a small pneumothorax. {extra}"])
    elif r < 0.85:
        lab = "absent"
        txt = rr.choice([f"{view} chest radiograph. No pneumothorax. {extra}",
                         f"No pneumothorax or pleural effusion identified. {extra}",
                         f"Chest tube in place; no residual pneumothorax. {extra}",
                         f"Previously seen {side} pneumothorax has resolved. {extra}",
                         f"{view} view. No evidence of pneumothorax. {extra}"])
    else:
        lab = "uncertain"
        txt = rr.choice([f"{view} chest radiograph. Skin fold projects over the {side} upper zone, difficult to exclude small pneumothorax. {extra}",
                         f"Questionable thin lucency along the {side} lateral chest wall, likely artifact; clinical correlation advised. {extra}"])
    reports.append((f"R{i:03d}", e["name"], txt)); rtruth.append((f"R{i:03d}", lab))
with open(KIT / "radiology_reports.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["report_id", "image_file", "report_text"]); w.writerows(reports)
with open(KIT / "radiology_reports_labeled_sample.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["report_id", "category"]); w.writerows(rtruth[:30])
with open(REF / "radiology_reports_truth.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["report_id", "category"]); w.writerows(rtruth)

# ---------------- text documents ----------------
OPS = json.load(open(HERE / "assets" / "operating_points.json"))
(KIT / "data_dictionary.md").write_text(f"""# Data dictionary

**Toy example.** This kit is a teaching demonstration, not a validation study. The images are a small public sample; the problems added
to the files, the reports, and the follow-up images are synthetic. Do not use any result for a clinical or purchasing decision.

## Files

| File | What it is |
|---|---|
| `images/img_NNNN.dcm` | Chest radiographs in DICOM format (one image per file) |
| `labels.csv` | Reference label for each image: whether a pneumothorax is present (`image_file`, `label`) |
| `model_scores.csv` | Output of four pretrained pneumothorax detectors, one column per model |
| `radiology_reports.csv` | Short free-text radiology reports for 200 of the images |
| `radiology_reports_labeled_sample.csv` | Correct report category for the first 30 reports |
| `eval_protocol.docx` | The evaluation protocol: scope, labels, metrics and the pilot bar |
| `radiology_notes.txt` | The request from the department |

## Fields

- `label`: `Pneumothorax` (air in the pleural space is present) or `No Pneumothorax`. Established by expert review for the source challenge.
- Model score columns: higher means the model is more confident that a pneumothorax is present. Each is meant to be a probability between 0 and 1,
  but the scales are **not comparable between models** (they are not calibrated), so compare models by ranking quality (AUC), not by raw values.
  All four are open-source DenseNet-121 models from the TorchXRayVision library, run on a CPU; they differ in training data:

| Column | Trained on | Published operating point (threshold) |
|---|---|---|
| `score_all` | Several public datasets combined (**primary model**) | {OPS["all"]:.4f} |
| `score_chex` | CheXpert only | {OPS["chex"]:.4f} |
| `score_nih` | NIH ChestX-ray14 only | {OPS["nih"]:.4f} |
| `score_pc` | PadChest only | {OPS["pc"]:.4f} |

- **Possible training overlap.** These images are believed to come from the NIH ChestX-ray14 data, which is part of the training data of
  `score_all` and `score_nih`. Treat those two models' numbers as optimistic. This is one reason the exercise is a demonstration only.
- DICOM headers follow the standard: `PatientID`, `PatientSex` (M/F), `PatientAge` (e.g. 062Y), `Modality` (CR = computed radiography),
  `ViewPosition` (PA = back-to-front, upright; AP = front-to-back, often portable/bedside), `PixelSpacing` (millimetres per pixel),
  `PhotometricInterpretation` (MONOCHROME2: low values are dark; MONOCHROME1: low values are bright), `SOPInstanceUID` (unique per image).
- Pixel values are 8-bit. Images are meant to be 1024 x 1024.

## Notes

- The images were de-identified before public release. Names, IDs and dates you see were assigned for this workshop. Any "patient"
  information that looks real in a few files was added synthetically as part of an exercise.
- Reports are synthetic and written from templates; they do not describe the specific images they are attached to.
- Blank in a field means "not recorded".
""")
(KIT / "radiology_notes.txt").write_text("""From: Radiology Informatics
To: Analyst
Re: Pneumothorax detector - is it ready for a pilot?

(Workshop toy example: a demonstration, not a validation. The images are a small public sample and the exercise is synthetic.)

We are considering a pilot of an open-source AI tool that scores chest radiographs for pneumothorax (air that has leaked into the space
around the lung). Before the department agrees, we need an independent look at how well it works on a sample of images. The sample is in
this folder. The detailed rules are in eval_protocol.docx. Our primary model is score_all; the other three columns are alternatives.

Please answer:

1. How well does the primary model separate images with a pneumothorax from those without, and how certain are you of that number?
2. Is the primary model really better than the three alternatives, or could the differences be chance? Compare them on the same patients.
3. Does the primary model work equally well for PA and AP views, men and women, and patients under and over 65? Be careful about what a small
   sample can and cannot show.
4. Using each model's published operating point, how does it do on sensitivity and specificity, and does any model meet the pilot bar in the protocol?
5. Separately, sort the 200 short radiology reports into present / absent / uncertain so we can compare the text with the tool.

The sample was assembled quickly from several sources and has not been checked. Please check it before using it, and tell us what you
changed. Do not use real patient data for any exercise: everything here is public or synthetic.
""")
from docx import Document
d = Document(); d.add_heading("Pneumothorax detector evaluation: protocol (synthetic toy example)", 1)
for h, body in [
    ("Status", "Teaching demonstration. Not a validation study; results must not inform any clinical or purchasing decision. The model training data may overlap with these images (see data_dictionary.md)."),
    ("Purpose", "Estimate how well pretrained pneumothorax detectors separate positive from negative frontal chest radiographs in a small sample, and whether the primary model is better than three alternatives."),
    ("Models", "Primary: score_all. Comparators: score_chex, score_nih, score_pc. All are used as supplied; nothing is trained or tuned on this sample."),
    ("Scope (inclusion)", "Adults (age 18 or older, age must be recorded and plausible) with a frontal chest radiograph (PA or AP view), 1024 x 1024 pixels, "
        "with a reference label and all four scores. Exclude images that cannot be decoded, are not chest radiographs, or are mirrored or otherwise unsafe to interpret. "
        "Exclusion rules are fixed before any model score is looked at."),
    ("Reference label", "The label in labels.csv, normalized to Pneumothorax / No Pneumothorax. If copies of the same image carry different labels, exclude all of them and report them."),
    ("Unit of analysis", "The patient. If a patient has several images, use their earliest study as the index image. Patient identity is the normalized patient ID."),
    ("No train/test split", "Every patient is a test patient: the models are pretrained and no threshold is tuned on this sample (see Operating point)."),
    ("Primary analysis", "Area under the ROC curve (AUC) of the primary model, with a 95% confidence interval from a patient-level bootstrap (at least 2000 resamples). Report the number of patients and positives."),
    ("Model comparison", "Compare the primary model with each alternative using the paired difference in AUC on the same patients (paired bootstrap, 95% interval). Do not compare separate intervals by eye."),
    ("Operating point", f"Each model's published operating point, fixed in advance: all {OPS['all']:.4f}, chex {OPS['chex']:.4f}, nih {OPS['nih']:.4f}, pc {OPS['pc']:.4f}. "
        "Report sensitivity and specificity at that threshold with 95% intervals. A threshold chosen after looking at these data gives optimistic results and must be labelled as such."),
    ("Subgroups", "For the primary model, compare PA vs AP, male vs female, and age under 65 vs 65 or older. Report counts (including positives), the difference in AUC with its 95% interval, and treat these as exploratory: "
        "they were not pre-registered and there are several of them."),
    ("Pilot bar", "A pilot would be considered if one threshold gives sensitivity of at least 0.85 and specificity of at least 0.80, with the lower end of each 95% interval no more than 0.10 below those values."),
    ("Reporting", "State every exclusion and fix with counts. State what could not be checked. Make no claim about clinical benefit."),
]:
    d.add_heading(h, 2); d.add_paragraph(body)
d.save(KIT / "eval_protocol.docx")

# ---------------- reference ----------------
meta = {e["name"]: dict(kind=e["kind"], source_path=e["source"], true_label=e["label"], true_scores=e["scores"],
                        problems=e["problems"], duplicate_of=e.get("dup_of"), fake_phi_text=e.get("fake_phi_text")) for e in entries}
(REF / "truth.json").write_text(json.dumps(dict(
    seed=2026, n_files=len(entries), siim_small_tgz_sha256=sha256(HERE / ".cache" / "siim_small.tgz"),
    planted=truth, files=meta,
    decoys_not_errors=["lossy JPEG compression (JPEG baseline transfer syntax) on untouched files", "varying PixelSpacing across scanners",
                       "AP and PA views mixed", "8-bit pixel depth", "bare-number PatientAge (e.g. '62') in most files (real quirk of the source data)",
                       "burned-in 'AP', 'PORTABLE', and L/R markers in real images (not PHI; possible shortcut features)",
                       "visible chest tubes / drains in many pneumothorax images (possible shortcut feature for a classifier)",
                       "a few real images are rotated, shifted or have black margins (the source data is heterogeneous; not planted)",
                       "pediatric patients (age under 18) in the sample: not an error, but excluded by the protocol"],
), indent=1))
print(len(entries), "files;", sum(1 for e in entries if e["problems"]), "with planted problems")
